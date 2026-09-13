#!/usr/bin/env python
"""
deck-ppsx · iterumcorp.org

Builds the read-only PowerPoint show (.ppsx) offered next to the PDF on pitch.html.

  · Every slide becomes one slide-sized picture, so nothing on it can be retyped.
  · The animated gifs and the videos are put back over their still frames as real media, in
    their exact place, so they play in the show (video: on click).
  · The file is saved as a *show* (opens straight into the slideshow) with a "password to
    modify" (PowerPoint opens it read-only unless the password is typed) and marked as final.
    This is the strongest protection an Office file has; it is not DRM — someone can still
    take screenshots or drag the pictures out. A file nobody can alter at all is a video.

Two sources:

  --pptx "…/Singular 26 Pitch.pptx"   PowerPoint itself renders the slides (Windows, PowerPoint
                                      installed), the gifs and videos come out of the .pptx with
                                      their geometry. Hidden slides are skipped. Videos over
                                      --video-max-mb are re-encoded to 720p/30 (needs ffmpeg:
                                      `pip install imageio-ffmpeg`). The same run also writes the
                                      two PDFs the web page needs (--pdf-out for download, --view-out
                                      for the viewer, with the transparent sprites left out so the
                                      animated gifs laid over it have no still frame behind them).
  --pdf assets/pitch/void-singularcorp.pdf
                                      Fallback without PowerPoint: the PDF pages are rendered and
                                      the clips listed in pitch.html (data-deck-page / data-deck-box)
                                      are laid over their stills.

    python scripts/deck-ppsx.py --pptx "C:/…/Singular 26 Pitch.pptx" --password "…"
    python scripts/deck-ppsx.py --pdf assets/pitch/void-singularcorp.pdf --password "…"
        [--out assets/pitch/void-singularcorp.ppsx] [--last N] [--width 1920] [--video-max-mb 60]

Needs python-pptx and Pillow; PyMuPDF for --pdf; pywin32 for --pptx:
    python -m pip install python-pptx pillow pymupdf pywin32 imageio-ffmpeg
"""
import argparse, base64, hashlib, io, os, re, secrets, shutil, struct, subprocess, sys, tempfile, zipfile

try:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu
except ImportError as e:
    sys.exit("missing dependency (%s): python -m pip install python-pptx pillow" % e)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPIN = 100000
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
CUSTOM_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="_MarkAsFinal"><vt:bool>true</vt:bool></property>'
              '</Properties>')


# ---------------------------------------------------------------------------------------------
# Source A: a .pptx, rendered by PowerPoint
# ---------------------------------------------------------------------------------------------
class PowerPoint:
    """One PowerPoint (COM) for the whole run. A PowerPoint the user already had open is used and
    left running; one started here is quit at the end. The deck is opened read-only for each job
    and closed without saving: PowerPoint writes only the first PDF of an opened presentation, so
    every export gets a fresh opening."""

    def __init__(self):
        try:
            import win32com.client
        except ImportError:
            sys.exit("--pptx needs PowerPoint and pywin32: python -m pip install pywin32")
        tl = subprocess.run(["tasklist"], capture_output=True, text=True).stdout.upper()
        self.was_running = "POWERPNT.EXE" in tl
        self.app = win32com.client.Dispatch("PowerPoint.Application")
        self.app.DisplayAlerts = 1               # ppAlertsNone

    def run(self, pptx, hide, job):
        """Open the deck, make the shapes in `hide` ({slide_no: {shape_id}}) invisible in memory,
        return job(pres), close."""
        path = os.path.abspath(pptx)
        try:
            pres = self.app.Presentations.Open(path, True, False, False)   # ReadOnly, Untitled, WithWindow
        except Exception:
            # Some PowerPoint builds refuse every windowless open ("could not open the file"): open it
            # in a window instead, minimised so it stays out of the way.
            pres = self.app.Presentations.Open(path, True, False, True)
            try:
                pres.Windows(1).WindowState = 2  # ppWindowMinimized
            except Exception:
                pass
        try:
            for i, ids in (hide or {}).items():
                for shp in pres.Slides(i).Shapes:
                    if shp.Id in ids:
                        shp.Visible = 0          # msoFalse
            return job(pres)
        finally:
            try:
                pres.Saved = True                # nothing to keep: no "save changes?" on the way out
            except Exception:
                pass
            pres.Close()

    def done(self):
        if not self.was_running:
            self.app.Quit()

    def export_pdf(self, pptx, path, hide=None):
        path = os.path.abspath(path)
        self.run(pptx, hide, lambda pres: pres.SaveCopyAs(path, 32))   # ppSaveAsPDF
        if not os.path.exists(path):
            sys.exit("PowerPoint did not write %s" % path)

    def render(self, pptx, outdir, width, hide=None):
        """Every visible slide to PNG, with the shapes in `hide` left out (the show's slide pictures:
        the gifs and videos are put back on top afterwards). Returns {slide_no: png_path}."""
        def job(pres):
            out = {}
            h = round(width * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
            for i in range(1, pres.Slides.Count + 1):
                s = pres.Slides(i)
                if s.SlideShowTransition.Hidden == -1:
                    continue
                p = os.path.normpath(os.path.join(outdir, "slide%02d.png" % i))
                s.Export(p, "PNG", width, h)
                out[i] = p
            return out
        return self.run(pptx, hide, job)


def media_from_pptx(pptx):
    """The animated gifs and embedded videos of every slide: {slide_no: [item]} with EMU geometry."""
    prs = Presentation(pptx)
    items = {}
    for i, slide in enumerate(prs.slides, 1):
        if slide._element.get("show") == "0":
            continue
        for sh in slide.shapes:
            box = (sh.left, sh.top, sh.width, sh.height)
            if None in box:
                continue
            vf = sh._element.find(".//" + NS_A + "videoFile")
            if vf is not None:
                rid = vf.get(NS_R + "link")
                rel = sh.part.rels.get(rid) if rid else None
                if rel is None or rel.is_external:
                    print("  ! slide %02d: linked (not embedded) video skipped: %s" % (i, rel.target_ref if rel else "?"))
                    continue
                poster = None
                try:
                    pf = sh.poster_frame
                    poster = pf.blob if pf is not None else None
                except Exception:
                    pass
                items.setdefault(i, []).append({"kind": "video", "blob": rel.target_part.blob, "box": box, "poster": poster,
                                                "mime": getattr(rel.target_part, "content_type", "video/mp4"), "id": sh.shape_id})
                continue
            try:
                img = sh.image
            except Exception:
                continue                          # not a picture, or a linked / SVG one
            if img.content_type != "image/gif":
                continue
            try:
                pil = Image.open(io.BytesIO(img.blob))
                frames = getattr(pil, "n_frames", 1)
                alpha = pil.convert("RGBA").getextrema()[3][0] < 255     # a sprite on a see-through ground
            except Exception:
                frames, alpha = 1, False
            if frames < 2:
                continue                          # a still gif is already in the render
            # The deck's own look for the picture: "crop to shape" (rounded corners…) and its outline.
            sp = sh._element.find(".//" + NS_P + "spPr")
            geom = sp.find(NS_A + "prstGeom") if sp is not None else None
            ln = sp.find(NS_A + "ln") if sp is not None else None
            items.setdefault(i, []).append({"kind": "gif", "blob": img.blob, "box": box, "id": sh.shape_id, "alpha": alpha,
                                            "crop": (sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom),
                                            "geom": geom, "ln": ln})
    return prs.slide_width, prs.slide_height, items


def dress(pic, geom, ln):
    """Give a new picture the source picture's shape geometry and outline (deep copies of the XML)."""
    import copy
    sp = pic._element.spPr
    if geom is not None and geom.get("prst") not in (None, "rect"):
        old = sp.find(NS_A + "prstGeom")
        new = copy.deepcopy(geom)
        if old is not None:
            sp.replace(old, new)
        else:
            sp.append(new)
    if ln is not None:
        old = sp.find(NS_A + "ln")
        if old is not None:
            sp.remove(old)
        sp.append(copy.deepcopy(ln))     # after prstGeom and any fill: the order CT_ShapeProperties wants


def ffmpeg_exe():
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except ImportError:
        return shutil.which("ffmpeg")


def shrink_video(blob, workdir, max_mb):
    """Re-encode a video that is too big to ship (720p, 30 fps, H.264, AAC). Returns a file path."""
    src = os.path.join(workdir, "video-src.mp4")
    with open(src, "wb") as f:
        f.write(blob)
    if len(blob) <= max_mb * 1e6:
        return src
    ff = ffmpeg_exe()
    if not ff:
        print("  ! video is %.0f MB and no ffmpeg found (pip install imageio-ffmpeg): embedding it as is" % (len(blob) / 1e6))
        return src
    dst = os.path.join(workdir, "video-720p.mp4")
    print("  re-encoding a %.0f MB video to 720p/30 …" % (len(blob) / 1e6))
    subprocess.run([ff, "-y", "-hide_banner", "-loglevel", "error", "-i", src, "-vf", "scale=1280:-2", "-r", "30",
                    "-c:v", "libx264", "-crf", "30", "-preset", "medium", "-pix_fmt", "yuv420p",
                    "-c:a", "aac", "-b:a", "128k", "-ac", "2", "-movflags", "+faststart", dst], check=True)
    print("  → %.1f MB" % (os.path.getsize(dst) / 1e6))
    return dst


# ---------------------------------------------------------------------------------------------
# Source B: the PDF plus the clips section of pitch.html
# ---------------------------------------------------------------------------------------------
def clips_from_html(path):
    """The figures of the clips section: {page: [{box(% of page), src, poster, video}]}."""
    html = open(path, encoding="utf-8").read()
    clips = {}
    for m in re.finditer(r'<figure\b[^>]*\bdata-deck-page="(\d+)"[^>]*\bdata-deck-box="([^"]+)"[^>]*>([\s\S]*?)</figure>', html):
        page, box, body = int(m.group(1)), [float(v) for v in m.group(2).split()], m.group(3)
        src = re.search(r'<(?:source|img)\b[^>]*\bsrc="([^"]+)"', body)
        poster = re.search(r'\bposter="([^"]+)"', body)
        if not src or len(box) != 4:
            continue
        clips.setdefault(page, []).append({"box": box, "src": src.group(1), "poster": poster.group(1) if poster else None,
                                           "video": bool(re.search(r"<video\b", body))})
    return clips


def render_pdf(pdf, outdir, width):
    try:
        import fitz
    except ImportError:
        sys.exit("--pdf needs PyMuPDF: python -m pip install pymupdf")
    doc = fitz.open(pdf)
    out = {}
    r = doc[0].rect
    for i, page in enumerate(doc, 1):
        pix = page.get_pixmap(matrix=fitz.Matrix(width / r.width, width / r.width), alpha=False)
        p = os.path.join(outdir, "slide%02d.png" % i)
        pix.save(p)
        out[i] = p
    return out, r.width / r.height


# ---------------------------------------------------------------------------------------------
# Build + protect
# ---------------------------------------------------------------------------------------------
def cover_crop(pic, iw, ih):
    """Crop a picture so it covers its box like object-fit: cover."""
    if iw / ih > pic.width / pic.height:
        cut = (1 - (pic.width / pic.height) / (iw / ih)) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - (iw / ih) / (pic.width / pic.height)) / 2
        pic.crop_top = pic.crop_bottom = cut


def set_background(slide, image_file):
    """The slide picture as the slide's own background (a:blipFill) rather than a picture shape:
    every viewer, phone ones included, paints a background first and under everything else, and
    nothing can select or drag it."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    _, rId = slide.part.get_or_add_image_part(image_file)
    bg = parse_xml('<p:bg %s><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="%s"/><a:srcRect/>'
                   '<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>' % (nsdecls("p", "a", "r"), rId))
    cSld = slide._element.find(NS_P + "cSld")
    old = cSld.find(NS_P + "bg")
    if old is not None:
        cSld.remove(old)
    cSld.insert(0, bg)


def build(renders, W, H, media, last, workdir):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    # One layout only. The default template carries ten more with 4:3 placeholders; nothing here
    # uses them and a phone viewer has less to misread without them.
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")
    for layout in list(prs.slide_layouts):
        if layout is not blank:
            prs.slide_layouts.remove(layout)
    n = 0
    for no in sorted(renders):
        if last and no > last:
            break
        img = Image.open(renders[no]).convert("RGB")
        buf = io.BytesIO(); img.save(buf, "JPEG", quality=86, optimize=True); buf.seek(0)
        slide = prs.slides.add_slide(blank)
        set_background(slide, buf)
        for m in media.get(no, []):
            left, top, width, height = (Emu(v) for v in m["box"])
            if m["kind"] == "video":
                poster = io.BytesIO(m["poster"]) if m.get("poster") else None
                slide.shapes.add_movie(m["path"], left, top, width, height, poster_frame_image=poster, mime_type=m.get("mime", "video/mp4"))
            else:
                pic = slide.shapes.add_picture(io.BytesIO(m["blob"]), left, top, width, height)
                if "crop" in m:                               # the deck's own crop
                    pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = m["crop"]
                else:                                         # a web clip over a still: cover the box
                    iw, ih = Image.open(io.BytesIO(m["blob"])).size
                    cover_crop(pic, iw, ih)
                dress(pic, m.get("geom"), m.get("ln"))       # rounded corners, outline: as in the deck
        n += 1
        print("  slide %02d%s" % (no, "  + %d clip(s)" % len(media[no]) if media.get(no) else ""))
    prs.core_properties.title = "Singular · Pitch"
    prs.core_properties.author = "Iterum Corporation"
    return prs, n


def verifier(password):
    """ECMA-376 write protection: SHA-512 of salt + UTF-16LE password, then SPIN rounds with a counter."""
    salt = secrets.token_bytes(16)
    h = hashlib.sha512(salt + password.encode("utf-16-le")).digest()
    for i in range(SPIN):
        h = hashlib.sha512(h + struct.pack("<I", i)).digest()
    return ('<p:modifyVerifier cryptProviderType="rsaAES" cryptAlgorithmClass="hash" cryptAlgorithmType="typeAny" '
            'cryptAlgorithmSid="14" spinCount="%d" saltData="%s" hashData="%s"/>'
            % (SPIN, base64.b64encode(salt).decode(), base64.b64encode(h).decode()))


def finish(tmp, out, password, show=True):
    """Turn the saved .pptx into a read-only .ppsx: modify password, mark as final, show content type."""
    ver = verifier(password) if password else ""
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/presentation.xml" and ver:
                s = data.decode("utf-8")
                s = s.replace("<p:extLst", ver + "<p:extLst", 1) if "<p:extLst" in s else s.replace("</p:presentation>", ver + "</p:presentation>")
                data = s.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                s = data.decode("utf-8")
                if show:                          # .ppsx opens straight into the slideshow; .pptx keeps the plain type
                    s = s.replace("presentationml.presentation.main+xml", "presentationml.slideshow.main+xml")
                if "docProps/custom.xml" not in names:
                    s = s.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
                data = s.encode("utf-8")
            elif item.filename == "_rels/.rels" and "docProps/custom.xml" not in names:
                s = data.decode("utf-8").replace("</Relationships>", '<Relationship Id="rIdCustomProps" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
                data = s.encode("utf-8")
            zout.writestr(item, data)
        if "docProps/custom.xml" not in names:
            zout.writestr("docProps/custom.xml", CUSTOM_XML)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group()
    src.add_argument("--pptx", help="the deck itself; PowerPoint renders the slides")
    src.add_argument("--pdf", help="the PDF export plus the clips of pitch.html (default: assets/pitch/void-singularcorp.pdf)")
    ap.add_argument("--html", default="pitch.html", help="where the clips come from in --pdf mode")
    ap.add_argument("--out", default="assets/pitch/void-singularcorp.ppsx")
    ap.add_argument("--pptx-out", default="assets/pitch/void-singularcorp.pptx", help="the same deck as a plain .pptx (--out with the extension swapped is a good choice for a private deck)")
    ap.add_argument("--no-pptx", action="store_true", help="build the .ppsx only")
    ap.add_argument("--password", default=None, help="password to modify (a random one is made and printed if omitted)")
    ap.add_argument("--last", type=int, default=None, help="stop after this slide")
    ap.add_argument("--width", type=int, default=1920, help="pixel width of the rendered slides")
    ap.add_argument("--video-max-mb", type=float, default=60, help="--pptx: videos bigger than this are re-encoded to 720p")
    ap.add_argument("--pdf-out", default="assets/pitch/void-singularcorp.pdf", help="--pptx: the PDF for download (the deck as it is)")
    ap.add_argument("--view-out", default="assets/pitch/void-singularcorp.view.pdf", help="--pptx: the PDF the web viewer draws (transparent sprites left out)")
    ap.add_argument("--no-pdfs", action="store_true", help="--pptx: build the show only")
    a = ap.parse_args()
    if not a.pptx and not a.pdf:
        a.pdf = "assets/pitch/void-singularcorp.pdf"
    out = os.path.join(ROOT, a.out)
    password = a.password if a.password is not None else secrets.token_urlsafe(9)
    work = tempfile.mkdtemp(prefix="deck-ppsx-")
    try:
        if a.pptx:
            pptx = os.path.abspath(a.pptx)
            W, H, media = media_from_pptx(pptx)
            hide = {n: {m["id"] for m in items} for n, items in media.items()}
            hide_view = {n: {m["id"] for m in items if m.get("alpha")} for n, items in media.items()}
            pp = PowerPoint()
            try:
                if not a.no_pdfs:
                    pdf_full, pdf_view = os.path.join(ROOT, a.pdf_out), os.path.join(ROOT, a.view_out)
                    print("exporting the PDFs with PowerPoint …")
                    pp.export_pdf(pptx, pdf_full)
                    pp.export_pdf(pptx, pdf_view, hide_view)
                    print("wrote %s (%.1f MB) and %s (%.1f MB, %d sprite(s) left out)" % (
                        a.pdf_out, os.path.getsize(pdf_full) / 1e6, a.view_out, os.path.getsize(pdf_view) / 1e6, sum(len(v) for v in hide_view.values())))
                print("rendering %s with PowerPoint …" % os.path.basename(pptx))
                renders = pp.render(pptx, work, a.width, hide)
            finally:
                pp.done()
            for items in media.values():
                for m in items:
                    if m["kind"] == "video":
                        m["path"] = shrink_video(m["blob"], work, a.video_max_mb)
                        del m["blob"]
            print("%d visible slide(s), media on: %s" % (len(renders), {k: len(v) for k, v in media.items()} or "none"))
        else:
            pdf, html = os.path.join(ROOT, a.pdf), os.path.join(ROOT, a.html)
            renders, ar = render_pdf(pdf, work, a.width)
            W, H = 12192000, round(12192000 / ar)
            media = {}
            for page, clips in clips_from_html(html).items():
                for c in clips:
                    x, y, w, h = c["box"]
                    box = (round(W * x / 100), round(H * y / 100), round(W * w / 100), round(H * h / 100))
                    path = os.path.join(ROOT, c["src"])
                    if c["video"]:
                        poster = open(os.path.join(ROOT, c["poster"]), "rb").read() if c["poster"] else None
                        media.setdefault(page, []).append({"kind": "video", "path": path, "box": box, "poster": poster, "mime": "video/mp4"})
                    else:
                        media.setdefault(page, []).append({"kind": "gif", "blob": open(path, "rb").read(), "box": box})
            print("%d page(s), clips from %s: %s" % (len(renders), a.html, {k: len(v) for k, v in media.items()} or "none"))
        prs, n = build(renders, W, H, media, a.last, work)
        tmp = os.path.join(work, "deck.pptx")
        prs.save(tmp)
        # Written beside the target and swapped in at the end, so a half-made file never replaces
        # the good one; if the old file is open in PowerPoint, say so instead of failing halfway.
        # The show (.ppsx) and, unless --no-pptx, the same deck as a plain .pptx for the viewers that do
        # not know what to do with a show (phones, mostly). Same pictures, same password, same "final".
        targets = [(out, a.out, True)]
        if not a.no_pptx:
            targets.append((os.path.join(ROOT, a.pptx_out), a.pptx_out, False))
        for path, label, show in targets:
            part = path + ".part"
            finish(tmp, part, password, show)
            try:
                os.replace(part, path)
            except PermissionError:
                os.remove(part)
                sys.exit("cannot replace %s: it is open in PowerPoint (or another program). Close it and run again." % label)
            print("wrote %s: %d slides, %.1f MB" % (label, n, os.path.getsize(path) / 1e6))
    finally:
        shutil.rmtree(work, ignore_errors=True)
    print("password to modify: %s" % (password if password else "(none)"))


if __name__ == "__main__":
    main()
