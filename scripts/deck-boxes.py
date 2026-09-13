#!/usr/bin/env python
"""
deck-boxes · iterumcorp.org

Prints, per slide, where the clips sit, as the attributes pitch.html expects on the figures of
its clips section: data-deck-page, data-deck-box ("left top width height", % of the slide) and,
when the deck crops a gif, data-deck-crop ("left right top bottom", fractions of the image).

From the deck itself (best: exact geometry, and the gifs can be pulled out of it):

    python scripts/deck-boxes.py "C:/…/Singular 26 Pitch.pptx" [--extract assets/img/deck]

    Lists every animated gif and every embedded video. --extract writes the gifs to that folder
    as deck-sNN-<name>.gif so pitch.html can reference them (existing files are not overwritten).

From a PDF export (fallback: every image box on every page, no crops):

    python scripts/deck-boxes.py assets/pitch/void-singularcorp.pdf [page ...]

Needs python-pptx + Pillow for a .pptx, PyMuPDF for a .pdf:  python -m pip install python-pptx pillow pymupdf
"""
import io, os, re, sys

if len(sys.argv) < 2:
    sys.exit(__doc__)
src = sys.argv[1]
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def from_pptx(path, extract):
    try:
        from pptx import Presentation
        from PIL import Image
    except ImportError:
        sys.exit("python -m pip install python-pptx pillow")
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    seen = {}
    for i, slide in enumerate(prs.slides, 1):
        if slide._element.get("show") == "0":
            print("\nslide %02d  (hidden, skipped)" % i)
            continue
        rows = []
        for sh in slide.shapes:
            if None in (sh.left, sh.top, sh.width, sh.height):
                continue
            box = "%.2f %.2f %.2f %.2f" % (sh.left / W * 100, sh.top / H * 100, sh.width / W * 100, sh.height / H * 100)
            vf = sh._element.find(".//" + NS_A + "videoFile")
            if vf is not None:
                rel = sh.part.rels.get(vf.get(NS_R + "link"))
                what = "video (linked: %s)" % rel.target_ref if rel is None or rel.is_external else "video %.0f MB" % (len(rel.target_part.blob) / 1e6)
                rows.append('  %-22s data-deck-page="%d" data-deck-box="%s"' % (what, i, box))
                continue
            try:
                img = sh.image
            except Exception:
                continue
            if img.content_type != "image/gif":
                continue
            try:
                frames = getattr(Image.open(io.BytesIO(img.blob)), "n_frames", 1)
            except Exception:
                frames = 1
            if frames < 2:
                continue
            crop = (sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom)
            attrs = 'data-deck-page="%d" data-deck-box="%s"' % (i, box)
            if any(abs(c) > 1e-4 for c in crop):
                attrs += ' data-deck-crop="%.4f %.4f %.4f %.4f"' % crop
            # "Crop to shape: rounded rectangle" → the corner radius as PowerPoint stores it (adj, 1/100000
            # of the shorter side; 16667 by default), and an outline → its width as % of the slide width
            # (the web viewer keeps the rendered outline and tucks the clip inside it).
            sp = sh._element.find(".//" + NS_P + "spPr")
            geom = sp.find(NS_A + "prstGeom") if sp is not None else None
            if geom is not None and geom.get("prst") == "roundRect":
                gd = geom.find(".//" + NS_A + "gd")
                adj = re.sub(r"\D", "", gd.get("fmla", "")) if gd is not None else ""
                attrs += ' data-deck-round="%s"' % (adj or "16667")
            ln = sp.find(NS_A + "ln") if sp is not None else None
            if ln is not None and ln.get("w") and ln.find(NS_A + "noFill") is None:
                attrs += ' data-deck-line="%.3f"' % (int(ln.get("w")) / W * 100)
            name = seen.get(img.sha1)
            if name is None:
                base = re.sub(r"[^a-z0-9]+", "-", (sh.name or "gif").lower()).strip("-") or "gif"
                name = "deck-s%02d-%s.gif" % (i, base)
                seen[img.sha1] = name
                if extract:
                    os.makedirs(extract, exist_ok=True)
                    dst = os.path.join(extract, name)
                    if not os.path.exists(dst):
                        with open(dst, "wb") as f:
                            f.write(img.blob)
            rows.append("  %-22s %s  ← %s (%d frames, %d KB)" % ("gif", attrs, name, frames, len(img.blob) // 1000))
        if rows:
            print("\nslide %02d" % i)
            print("\n".join(rows))


def from_pdf(path, pages):
    try:
        import fitz
    except ImportError:
        sys.exit("python -m pip install pymupdf")
    doc = fitz.open(path)
    only = {int(p) for p in pages}
    for i, page in enumerate(doc, 1):
        if only and i not in only:
            continue
        w, h = page.rect.width, page.rect.height
        text = " | ".join(page.get_text().split())[:70]
        print("\npage %02d  %.0fx%.0f  %s" % (i, w, h, text))
        for im in page.get_image_info(xrefs=True):
            x0, y0, x1, y1 = im["bbox"]
            if (x1 - x0) * (y1 - y0) < w * h * 0.02:   # icons and ornaments
                continue
            box = "%.2f %.2f %.2f %.2f" % (x0 / w * 100, y0 / h * 100, (x1 - x0) / w * 100, (y1 - y0) / h * 100)
            print('  %5dx%-5d  data-deck-page="%d" data-deck-box="%s"' % (im["width"], im["height"], i, box))


if src.lower().endswith(".pptx"):
    extract = sys.argv[sys.argv.index("--extract") + 1] if "--extract" in sys.argv else None
    from_pptx(src, extract)
else:
    from_pdf(src, [a for a in sys.argv[2:] if a.isdigit()])
